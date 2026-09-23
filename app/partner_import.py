"""Adapters for the supplied IEK/Systeme reports; the engine contract is unchanged.

Header inspection is bounded. Full worksheets are read only on explicit import.
No client identity, stockout interval, or current inventory is inferred.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from itertools import islice
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

BRANDS = {"IEK": "ИЭК", "SYSTEME": "Systeme Electric"}
WAREHOUSE = "Все склады (сводный отчёт)"
KINDS = {"sales": "Месячные продажи", "movements": "Динамика продаж",
         "stock": "Месячные остатки", "transit": "Поставки в пути",
         "overview": "Продажи, остатки и поставки", "moq": "MOQ / кратность",
         "season": "Сезонность (справочно)"}
MONTHS = {name: i for i, name in enumerate(
    ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"], 1)}
# Two supplied IEK reports omit the brand in the original filename.
IEK_FILENAMES = {"динамика продаж_2025-2026.xlsx",
                 "ежемесячные продажи в количественном выражении за последние 2 года.xlsx"}


def text(value):
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (float, np.floating)) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def norm(value):
    return text(value).casefold().replace("ё", "е")


def month(value):
    label = norm(value)
    match = re.search(r"20\d{2}", label)
    if match and label[:3] in MONTHS:
        return pd.Timestamp(int(match[0]), MONTHS[label[:3]], 1)
    return None


@dataclass(frozen=True)
class Report:
    filename: str
    sheet: str
    brand: str
    kind: str
    header_row: int
    headers: tuple[str, ...]


def _kind(headers):
    h = set(headers)
    if {"дата", "документ", "код", "количество", "склад"} <= h:
        return "movements"
    if "код 1с" in h and "свободный остаток" in h:
        return "overview"
    if "код 1с" in h and any("поступление до" in v for v in h):
        return "transit"
    if any(month(v) is not None for v in h) and "номенклатура.код" in h:
        return "stock" if h & {"ед.", "ед.изм"} else "sales"
    if h & {"код 1с", "номенклатура.код"} and h & {"кратность", "мин. разр. к отгр."}:
        return "moq"
    if {"год", "янв", "дек"} <= h:
        return "season"
    return None


def inspect_files(files: tuple) -> list[Report]:
    reports = []
    for filename, content in files:
        if Path(filename).suffix.lower() != ".xlsx":
            raise ValueError("Отчёты партнёра загружаются в XLSX. Подготовленные CSV загружайте отдельно.")
        name = norm(filename)
        brand = "SYSTEME" if any(s in name for s in ("system", "syseme")) else "IEK" if any(s in name for s in ("иэк", "iek")) else None
        if name in IEK_FILENAMES:
            brand = "IEK"
        if not brand:
            raise ValueError(f"{filename}: укажите ИЭК или Systeme в имени файла, чтобы определить поставщика.")
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            found = False
            try:
                for sheet in book:
                    sheet.reset_dimensions()  # Some 1C exports declare only H171605 as their dimension.
                    for rownum, row in enumerate(islice(sheet.iter_rows(values_only=True), 12), 1):
                        headers = tuple(norm(v) for v in row)
                        kind = _kind(headers)
                        if kind:
                            reports.append(Report(filename, sheet.title, brand, kind, rownum, headers))
                            found = True
                            break
            finally:
                book.close()
            if not found:
                raise ValueError(f"{filename}: не распознаны заголовки отчёта в первых 12 строках.")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"{filename}: не удалось открыть XLSX. Проверьте, что файл не повреждён и не защищён паролем.") from exc
    for brand in {r.brand for r in reports}:
        for kind in KINDS.keys() - {"season"}:
            if sum(r.brand == brand and r.kind == kind for r in reports) > 1:
                raise ValueError(f"{BRANDS[brand]}: несколько отчётов типа «{KINDS[kind]}». Оставьте одну версию.")
    return reports


def missing_reports(reports):
    missing = []
    for brand in sorted({r.brand for r in reports}):
        kinds = {r.kind for r in reports if r.brand == brand}
        if not kinds & {"sales", "movements", "overview"}:
            missing.append(f"{BRANDS[brand]}: продажи (месячный отчёт или динамика)")
        if not kinds & {"stock", "overview"}:
            missing.append(f"{BRANDS[brand]}: остатки")
    return missing


def _read(report, contents):
    book = load_workbook(io.BytesIO(contents[report.filename]), read_only=True, data_only=True)
    try:
        sheet = book[report.sheet]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(min_row=report.header_row + 1, max_col=len(report.headers), values_only=True)
        # Empty/duplicate unnamed columns are retained by position until headers are selected.
        frame = pd.DataFrame(rows, columns=[v or f"_empty_{i}" for i, v in enumerate(report.headers)])
    finally:
        book.close()
    code = next((c for c in ("код 1с", "номенклатура.код", "код") if c in frame), None)
    if code is None:
        raise ValueError(f"{report.filename}: нет кода товара.")
    frame["sku"] = frame[code].map(text)
    frame = frame.loc[frame.sku.ne("")].copy()
    if report.kind == "transit":
        meaningful_columns = [column for column in frame if not column.startswith("_empty_")]
        normalized = frame[meaningful_columns].copy()
        for column in meaningful_columns:
            normalized[column] = normalized[column].map(text)
        duplicate_rows = normalized.duplicated(keep="first")
        original_count = len(frame)
        frame = frame.loc[~duplicate_rows].copy()
        frame.attrs["exact_duplicate_rows_removed"] = original_count - len(frame)
    if report.kind == "moq" and frame.sku.duplicated().any():
        article = next((column for column in ("артикул поставщика", "артикул иэк", "артикул") if column in frame), None)
        constraint = next((column for column in ("кратность", "мин. разр. к отгр.") if column in frame), None)
        semantic_columns = [column for column in (article, constraint) if column]
        semantic = frame[["sku", *semantic_columns]].copy()
        for column in semantic_columns:
            semantic[column] = semantic[column].map(text)
        conflicting = semantic.groupby("sku", sort=False)[semantic_columns].nunique(dropna=False).gt(1).any(axis=1)
        if not conflicting.any():
            original_count = len(frame)
            frame = frame.drop_duplicates(subset=["sku"], keep="first").copy()
            frame.attrs["duplicate_moq_rows_removed"] = original_count - len(frame)
    if report.kind != "movements" and frame.sku.duplicated().any():
        raise ValueError(f"{report.filename}: код товара повторяется; уточните структуру, чтобы не удвоить объём.")
    return frame


def _numbers(values, label, nonnegative=False):
    raw = values.map(text).str.replace(" ", "", regex=False).str.replace(",", ".", regex=False)
    result = pd.to_numeric(raw.replace("", "0"), errors="coerce")
    bad = ~np.isfinite(result) | ((result < 0) if nonnegative else False)
    if bad.any():
        raise ValueError(f"{label}: нечисловое или недопустимое количество ({int(bad.sum())} строк). Исправьте исходный файл.")
    return result.astype(float)


def _monthly_sales(frame, today):
    parts = []
    for col in frame:
        date = month(col)
        if date is not None and date.to_period("M") < today.to_period("M"):
            parts.append(pd.DataFrame({"sku": frame.sku, "date": date,
                                       "qty": _numbers(frame[col], col)}))
    if not parts:
        raise ValueError("В отчёте нет завершённых месяцев до даты расчёта.")
    return pd.concat(parts, ignore_index=True)


def convert_files(files, reports, settings, today):
    """Return canonical raw tables + warnings. Must still pass engine.io.load_data."""
    today = pd.Timestamp(today)
    missing = missing_reports(reports)
    if missing:
        raise ValueError("Файлы распознаны. Для расчёта добавьте: " + "; ".join(missing))
    if not settings.get("confirmed"):
        raise ValueError("Подтвердите настройки и ограничения импорта перед расчётом.")
    from engine.io import REQUIRED_COLUMNS
    tables = {key: pd.DataFrame(columns=cols) for key, cols in REQUIRED_COLUMNS.items()}
    contents = dict(files)
    if len(contents) != len(files):
        raise ValueError("Имена загруженных файлов повторяются. Оставьте одну версию каждого файла.")
    warnings = [
        "Импорт: сводный расчёт по всем складам. Разбивку по складам из месячных отчётов восстановить нельзя.",
        "Нет клиентов и точных периодов отсутствия: разовые клиентские заказы и упущенный спрос на этих данных не подтверждаются. BOM не предоставлен.",
        "Пустые количества внутри отчётов приняты за 0. Текущий незавершённый месяц не используется для обучения.",
        "Единицы покупки и хранения не пересчитываются (например, бухты → метры). Проверьте кратность перед утверждением.",
    ]
    for brand in sorted({r.brand for r in reports}):
        chosen = {r.kind: r for r in reports if r.brand == brand and r.kind != "season"}
        config = settings.get(brand, {})
        lead, cycle = config.get("lead_time_days"), config.get("order_cycle_days")
        if not all(isinstance(v, (int, float)) and np.isfinite(v) and 1 <= v <= 365 for v in (lead, cycle)):
            raise ValueError(f"{BRANDS[brand]}: задайте срок поставки и цикл заказа от 1 до 365 дней.")
        # Prefer the complete monthly report, never add detail lines on top of monthly totals.
        sales_kind = next(k for k in ("sales", "movements", "overview") if k in chosen)
        stock_kind = "overview" if "overview" in chosen else "stock"
        needed = {sales_kind, stock_kind} | ({"moq", "transit", "stock"} & chosen.keys())
        frames = {k: _read(chosen[k], contents) for k in needed}
        duplicate_transit_rows = sum(
            int(frame.attrs.get("exact_duplicate_rows_removed", 0)) for frame in frames.values()
        )
        if duplicate_transit_rows == 1:
            warnings.append(
                f"{BRANDS[brand]}: удалена 1 полностью совпадающая строка поставки, чтобы не удвоить объём."
            )
        elif duplicate_transit_rows > 1:
            warnings.append(
                f"{BRANDS[brand]}: удалено {duplicate_transit_rows} полностью совпадающих строк поставки, "
                "чтобы не удвоить объём."
            )
        duplicate_moq_rows = sum(
            int(frame.attrs.get("duplicate_moq_rows_removed", 0)) for frame in frames.values()
        )
        if duplicate_moq_rows == 1:
            warnings.append(
                f"{BRANDS[brand]}: удалена 1 повторная строка MOQ с теми же артикулом и ограничением."
            )
        elif duplicate_moq_rows > 1:
            warnings.append(
                f"{BRANDS[brand]}: удалено {duplicate_moq_rows} повторных строк MOQ "
                "с теми же артикулами и ограничениями."
            )
        if "movements" in chosen and sales_kind != "movements":
            warnings.append(f"{BRANDS[brand]}: использованы месячные продажи; динамика не прибавляется к ним.")
        sf = frames[sales_kind]
        if sales_kind == "movements":
            dates = pd.to_datetime(sf["дата"].map(text), dayfirst=True, errors="coerce")
            if dates.isna().any():
                raise ValueError(f"{BRANDS[brand]}: в динамике есть некорректные даты.")
            # Source movements record outgoing sales as negative; returns are positive.
            qty = -_numbers(sf["количество"], "Динамика: количество")
            documents = sf["документ"].map(norm)
            if not documents.str.contains("расходная накладная|возврат", regex=True).all():
                raise ValueError(f"{BRANDS[brand]}: неизвестный тип документа в динамике. Загрузите месячные продажи или уточните правила знака.")
            if ((documents.str.contains("расходная накладная") & (qty < 0))).any():
                raise ValueError(f"{BRANDS[brand]}: неожиданный знак расходной накладной. Проверьте выгрузку.")
            sales = pd.DataFrame({"sku": sf.sku, "date": dates.dt.to_period("M").dt.to_timestamp(), "qty": qty})
            sales = sales[sales.date < today.to_period("M").start_time]
            sales = sales.groupby(["sku", "date"], as_index=False).qty.sum()
            warnings.append(f"{BRANDS[brand]}: знак движений обращён; строки агрегированы помесячно, без выдуманных идентификаторов клиентов.")
        else:
            sales = _monthly_sales(sf, today)
        if sales.empty or not (sales.qty > 0).any():
            raise ValueError(f"{BRANDS[brand]}: нет положительных продаж в завершённых месяцах.")
        sales = sales.assign(warehouse=WAREHOUSE, client_id="UNKNOWN", price=0.)
        # One aggregate per SKU/month with the same unknown identity cannot be
        # flagged as an irregular customer by the current (>=3 lines) detector.
        tables["sales"] = pd.concat([tables["sales"], sales], ignore_index=True)

        inventory = frames[stock_kind]
        if stock_kind == "overview":
            date_match = re.search(r"(\d{2}\.\d{2}\.20\d{2})", chosen[stock_kind].filename)
            if not date_match:
                raise ValueError("Для сводного отчёта Systeme укажите дату среза ДД.ММ.ГГГГ в имени файла.")
            snapshot = pd.to_datetime(date_match[1], dayfirst=True)
            if snapshot > today:
                raise ValueError("Дата остатков позже даты расчёта. Выберите дату не раньше среза.")
            stock_col = "свободный остаток"
        else:
            cols = [(month(c), c) for c in inventory if month(c) is not None and month(c) <= today]
            if not cols:
                raise ValueError(f"{BRANDS[brand]}: нет остатков на дату расчёта или раньше.")
            snapshot, stock_col = max(cols)
        stock = pd.DataFrame({"sku": inventory.sku, "on_hand": _numbers(inventory[stock_col], stock_col, True), "warehouse": WAREHOUSE})
        tables["stock"] = pd.concat([tables["stock"], stock], ignore_index=True)
        warnings.append(f"{BRANDS[brand]}: остатки из «{stock_col}», срез {snapshot:%d.%m.%Y}; движения после среза не восстановлены.")

        product_parts = []
        for kind, frame in frames.items():
            name_col = next((c for c in ("номенклатура", "наименование") if c in frame), None)
            article_col = next((c for c in ("артикул поставщика", "артикул иэк", "артикул") if c in frame), None)
            unit_col = next((c for c in ("ед.", "ед.изм") if c in frame), None)
            part = pd.DataFrame({"sku": frame.sku, "name": frame[name_col].map(text) if name_col else "",
                                 "article": frame[article_col].map(text) if article_col else "",
                                 "unit": frame[unit_col].map(text) if unit_col else ""})
            part["category"] = frame["категория 2026"].map(text) if "категория 2026" in frame else ""
            product_parts.append(part)
        products = pd.concat(product_parts).replace("", pd.NA).groupby("sku", sort=False).first().reset_index()
        products["name"] = products.name.fillna(products.sku)
        products["article"] = products.article.fillna(products.sku)
        products["unit"] = products.unit.fillna("ед.")
        products["category"] = BRANDS[brand] + " / " + products.category.fillna("Без категории")
        products = products.assign(supplier_id=brand, pack_size=1., moq=0., import_note="Сводный импорт; нет клиентов, stockout и BOM.")
        if "moq" in frames:
            mf = frames["moq"].set_index("sku")
            column = "кратность" if "кратность" in mf else "мин. разр. к отгр."
            values = _numbers(mf[column], column, True)
            if (values % 1 != 0).any():
                raise ValueError("Дробная кратность/MOQ не поддерживается целочисленным заказом. Уточните единицы измерения.")
            field = "pack_size" if column == "кратность" else "moq"
            products[field] = products.sku.map(values).fillna(1. if field == "pack_size" else 0.).clip(lower=1. if field == "pack_size" else 0.)
            uncovered = (~products.sku.isin(mf.index)).sum()
            if uncovered:
                warnings.append(f"{BRANDS[brand]}: {uncovered} товаров отсутствуют в MOQ; для них кратность 1, минимум 0.")
        else:
            warnings.append(f"{BRANDS[brand]}: MOQ не загружен; приняты кратность 1, минимум 0.")
        missing_stock = products.sku.nunique() - products.sku.isin(stock.sku).sum()
        if missing_stock:
            warnings.append(f"{BRANDS[brand]}: {missing_stock} товаров без остатка не участвуют в заказе; нулевой остаток им не подставляется.")
        tables["products"] = pd.concat([tables["products"], products], ignore_index=True)
        tables["suppliers"] = pd.concat([tables["suppliers"], pd.DataFrame([dict(
            supplier_id=brand, name=BRANDS[brand], lead_time_days=lead, order_cycle_days=cycle)])], ignore_index=True)
        for kind in ("transit", "overview"):
            if kind not in frames:
                continue
            frame = frames[kind]
            for col in frame:
                if "поступление до" not in col and not col.startswith("сэ в пути"):
                    continue
                match = re.search(r"(\d{2}\.\d{2}(?:\.20\d{2})?)", col.split("поступление до")[-1])
                if not match:
                    raise ValueError(f"Не распознана дата поставки: {col}")
                date_text = match[1]
                if len(date_text) == 5:
                    year = re.findall(r"20\d{2}", chosen[kind].filename)
                    if not year:
                        raise ValueError("Укажите год отчёта о поставках в имени файла.")
                    date_text += "." + year[-1]
                eta = pd.to_datetime(date_text, format="%d.%m.%Y", errors="coerce")
                if pd.isna(eta):
                    raise ValueError(f"Некорректная дата поставки: {date_text}")
                part = pd.DataFrame({"sku": frame.sku, "warehouse": WAREHOUSE,
                                     "qty": _numbers(frame[col], col, True), "eta": eta})
                tables["in_transit"] = pd.concat([tables["in_transit"], part[part.qty > 0]], ignore_index=True)
    if tables["products"].sku.duplicated().any():
        raise ValueError("Один код 1С относится к двум поставщикам. Уточните справочник; автоматическое объединение небезопасно.")
    if any(r.kind == "season" for r in reports):
        warnings.append("Листы сезонности распознаны справочно. Движок рассчитывает сезонность по истории количества; готовые коэффициенты повторно не применяются.")
    return tables, warnings
