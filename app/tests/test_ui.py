"""UI contract tests. Fixtures below are test data, never a production fallback."""

import io
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from app import backend, orders
from app.charts import demand_chart
from app.checks import grouping, run_checks


@pytest.fixture
def result():
    rows = []
    for sku, supplier, warehouse, quantity in [
        ("DEMO-ONEOFF", "S1", "Алматы", 100),
        ("DEMO-ONEOFF", "S1", "Астана", 50),
        ("DEMO-STOCKOUT", "S2", "Алматы", 40),
        ("DEMO-TRANSIT", "S2", "Алматы", 0),
    ]:
        rows.append(dict(supplier_id=supplier, supplier_name="Поставщик / " + supplier,
                         sku=sku, name="Товар " + sku, category="Кабель", warehouse=warehouse,
                         on_hand=10., in_transit=20., forecast_horizon=100., recommended_qty=quantity,
                         urgency="critical", confidence="high", explanation="Прогноз и страховой запас.",
                         transit_overdue=False, days_of_cover=5., safety_stock=30., need_raw=quantity,
                         growth_pct=0., flags=""))
    return pd.DataFrame(rows)


@pytest.fixture
def connected(monkeypatch, result, tmp_path):
    tables = {
        "products": result[["sku", "name", "category", "supplier_id"]].drop_duplicates("sku"),
        "stock": result[["sku", "warehouse", "on_hand"]],
        "growth": pd.DataFrame(columns=["category", "growth_pct"]),
    }
    monkeypatch.setattr(backend, "demo_files", lambda: (("sales.csv", b"test-fixture"),))
    monkeypatch.setattr(backend, "load_files", lambda files, today=None: (tables, []))
    monkeypatch.setattr(backend, "calculate", lambda data, params: result.copy())
    monkeypatch.setattr(backend, "series", lambda *args: pd.DataFrame({
        "month": ["2026-07-01", "2026-08-01", "2026-09-01"],
        "raw": [100., 150., np.nan], "clean": [100., 110., np.nan],
        "forecast": [np.nan, np.nan, 120.], "oneoff": [0, 40, 0], "stockout_days": [0, 3, 0],
    }))
    real_approve = orders.approve
    monkeypatch.setattr(orders, "approve", lambda frame, author, params: real_approve(frame, author, params, tmp_path))
    return tmp_path


def test_approval_export_and_immutable_history(result, tmp_path):
    original = result[result.supplier_id == "S1"].reset_index(drop=True)
    edits = original.copy()
    edits.loc[0, "recommended_qty"] = 75
    edits.loc[1, "recommended_qty"] = 0
    draft = orders.edited_order(original, edits)
    approved = orders.approve(draft, "Амир", {"today": "2026-09-23"}, tmp_path)
    assert approved["status"] == "approved"
    assert approved["items"][0]["original_qty"] == 100
    assert approved["items"][0]["manually_changed"] is True
    assert approved["fingerprint"] == orders.fingerprint(draft)
    orders.approve(draft, "Амир", {}, tmp_path)
    assert len(list(tmp_path.glob("*.json"))) == 2
    assert json.loads(next(tmp_path.glob("*.json")).read_text(encoding="utf-8"))["author"] == "Амир"
    exported = pd.read_csv(io.BytesIO(orders.csv_bytes(draft)), sep=";")
    assert list(exported["Количество"]) == [75]
    assert "100 → 75" in exported.iloc[0]["Обоснование"]
    assert orders.csv_bytes(draft).startswith(b"\xef\xbb\xbf")
    workbook = load_workbook(io.BytesIO(orders.xlsx_bytes(draft)))
    assert workbook.active["D2"].value == 75
    assert workbook.active.max_row == 2


@pytest.mark.parametrize("quantity", [-1, None, float("nan"), float("inf"), 1.5, 1e10])
def test_invalid_quantity_rejected(result, quantity):
    edits = result.copy().astype({"recommended_qty": float})
    edits.loc[0, "recommended_qty"] = quantity
    with pytest.raises(ValueError, match="целым"):
        orders.edited_order(result, edits)


def test_blank_author_and_empty_order_rejected(result, tmp_path):
    frame = result[result.supplier_id == "S1"].copy()
    with pytest.raises(ValueError, match="имя"):
        orders.approve(frame, "  ", {}, tmp_path)
    frame["recommended_qty"] = 0
    with pytest.raises(ValueError, match="положительным"):
        orders.approve(frame, "Амир", {}, tmp_path)
    assert not list(tmp_path.glob("*.json"))


def test_export_sheet_names_and_formula_protection(result):
    result.loc[0, "name"] = '=HYPERLINK("bad")'
    result["supplier_name"] = "Имя/с[символами]" * 5
    workbook = load_workbook(io.BytesIO(orders.xlsx_bytes(result)))
    assert len(workbook.sheetnames) == 2
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert workbook.worksheets[0]["B2"].data_type == "s"
    assert workbook.worksheets[0]["B2"].value.startswith("'=")


def test_result_validation(result):
    backend.validate_result(result)
    with pytest.raises(ValueError, match="отсутствуют"):
        backend.validate_result(result.drop(columns="sku"))
    with pytest.raises(ValueError, match="повторяется"):
        backend.validate_result(pd.concat([result, result]))
    result.loc[0, "recommended_qty"] = -1
    with pytest.raises(ValueError, match="Некорректные"):
        backend.validate_result(result)


def test_main_edit_approve_invalidate_and_recalculate(connected):
    app = AppTest.from_file(str(ROOT / "app/streamlit_app.py"), default_timeout=20).run()
    def click(label):
        next(button for button in app.button if button.label == label).click().run()
    assert not app.exception
    click("Рассчитать")
    assert not app.exception
    assert len(app.get("download_button")) == 0
    click("Утвердить заказ · Поставщик / S1")
    assert any("Укажите имя" in warning.value for warning in app.warning)
    app.text_input(key="approver").set_value("Амир").run()
    click("Утвердить заказ · Поставщик / S1")
    assert not app.exception
    assert len(app.get("download_button")) == 4
    assert len(list(connected.glob("*.json"))) == 1
    # Streamlit's test API exposes data_editor changes through widget session state.
    editor_key = "editor_" + hashlib.sha256(b"S1").hexdigest()[:12] + "_1"
    app.session_state[editor_key] = {"edited_rows": {0: {"recommended_qty": 77}}, "added_rows": [], "deleted_rows": []}
    app.run()
    assert not app.exception
    assert len(app.get("download_button")) == 0
    assert any("Изменено вручную" in caption.value for caption in app.caption)
    # AppTest does not model the editor's frontend; resend its browser state
    # with the subsequent button event, as the browser does.
    app.session_state[editor_key] = {"edited_rows": {0: {"recommended_qty": 77}}, "added_rows": [], "deleted_rows": []}
    click("Утвердить заказ · Поставщик / S1")
    assert app.session_state["approvals"]["S1"]["items"][0]["recommended_qty"] == 77
    assert len(list(connected.glob("*.json"))) == 2
    assert len(app.get("download_button")) == 4
    app.switch_page("pages/2_Проверки.py").run()
    assert not app.exception
    app.switch_page("streamlit_app.py").run()
    assert not app.exception
    assert len(app.get("download_button")) == 4
    assert any("77" in caption.value and "Изменено вручную" in caption.value for caption in app.caption)
    next(box for box in app.selectbox if box.label == "Склад").select("Алматы").run()
    assert len(app.get("download_button")) == 0
    assert next(button for button in app.button if button.label == "Утвердить заказ · Поставщик / S1").disabled
    click("Рассчитать")
    assert not app.session_state["approvals"]


def test_missing_backend_is_user_visible(monkeypatch):
    def unavailable():
        raise backend.BackendUnavailable("Демонстрационные данные ещё не добавлены командой.")
    monkeypatch.setattr(backend, "demo_files", unavailable)
    app = AppTest.from_file(str(ROOT / "app/streamlit_app.py")).run()
    assert not app.exception
    assert app.button[0].disabled
    assert any("ещё не добавлены" in warning.value for warning in app.warning)


def test_checks_page_without_calculation():
    app = AppTest.from_file(str(ROOT / "app/pages/2_Проверки.py")).run()
    assert not app.exception
    assert any("Сначала" in element.value for element in app.info)


def test_grouping_rejects_missing_supplier_and_explanation(result):
    result.loc[0, "supplier_id"] = None
    result.loc[1, "explanation"] = " "
    rows = grouping({}, {}, lambda data, params: result)
    assert any(row["Результат"] == "❌" for row in rows)


def test_checks_never_mark_missing_data_passed(monkeypatch, result):
    import app.checks as checks
    monkeypatch.setattr(checks, "engine_function", lambda *args: lambda data, params: result)
    report = run_checks({}, {"today": "2026-09-23"})
    assert len(report) == 5
    assert all(check.status != "passed" for check in report[:4])
    assert report[4].status == "passed"


def test_chart_handles_gaps_and_annotations():
    chart = demand_chart(pd.DataFrame({"month": ["2026-01", "2026-02"],
                                     "raw": [10, None], "clean": [12, None], "forecast": [None, 13],
                                     "oneoff": [5, 0], "stockout_days": [3, 0]}))
    assert len(chart.data) == 4
    assert len(chart.layout.shapes) == 1


def test_failed_recalculation_blocks_previous_approval(connected, monkeypatch):
    app = AppTest.from_file(str(ROOT / "app/streamlit_app.py")).run()
    next(button for button in app.button if button.label == "Рассчитать").click().run()
    app.text_input(key="approver").set_value("Амир").run()
    next(button for button in app.button if button.label.endswith("Поставщик / S1")).click().run()
    assert app.get("download_button")
    def fail(*args):
        raise ValueError("Некорректные входные данные")
    monkeypatch.setattr(backend, "calculate", fail)
    next(button for button in app.button if button.label == "Рассчитать").click().run()
    assert not app.exception
    assert not app.get("download_button")
    assert any("Некорректные входные данные" in warning.value for warning in app.warning)


def test_uploaded_tables_use_engine_loader(monkeypatch):
    backend.load_files.clear()
    seen = []
    def loader(path):
        seen.extend(sorted(file.name for file in path.glob("*.csv")))
        return {"stock": pd.DataFrame(columns=["warehouse"]),
                "products": pd.DataFrame(columns=["category"])}, ["Тестовое предупреждение"]
    monkeypatch.setattr(backend, "engine_function", lambda *args: loader)
    files = tuple((f"{name}.csv", b"header\n") for name in ["sales", "stock", "products", "suppliers"])
    _, warnings = backend.load_files(files)
    assert seen == ["products.csv", "sales.csv", "stock.csv", "suppliers.csv"]
    assert warnings == ["Тестовое предупреждение"]
    with pytest.raises(ValueError, match="дважды"):
        backend.load_files(files + (files[0],))
    with pytest.raises(ValueError, match="Не хватает"):
        backend.load_files(files[:1])
    backend.load_files.clear()
