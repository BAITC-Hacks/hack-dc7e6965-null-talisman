"""Small synthetic workbooks with the partner headers; no partner records in Git."""
import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from app import backend, orders
from app.partner_import import WAREHOUSE, convert_files, inspect_files, missing_reports

APP = Path(__file__).resolve().parents[1]


def workbook(rows, sheet="Лист_1"):
    book = Workbook()
    book.active.title = sheet
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def settings(**overrides):
    return {"confirmed": True, "IEK": {"lead_time_days": 14, "order_cycle_days": 14},
            "SYSTEME": {"lead_time_days": 30, "order_cycle_days": 14}, **overrides}


@pytest.fixture
def iek_files():
    months = [f"{m} 2026" for m in ("янв.", "февр.", "март", "апр.", "май", "июнь", "июль", "авг.", "сент.")]
    return (
        ("Ежемесячные продажи в количественном выражении за последние 2 года.xlsx", workbook([
            ["Номенклатура", "Номенклатура.Код", *months, "Итого"],
            [None, None, *["Количество"] * 10],
            ["Кабель тестовый", "001_", *[100] * 8, 90000, 90800],
        ])),
        ("Ежемесячные остатки ИЭК.xlsx", workbook([
            ["Номенклатура", "Ед.", "Номенклатура.Код", "авг. 2026", "сент. 2026", "Итого"],
            [None, None, None, "Количество", "Количество", "Количество"],
            [None, None, None, "нач. остаток", "нач. остаток", "нач. остаток"],
            ["Кабель тестовый", "м", "001_", 30, 2, 1000],
        ])),
        ("MOQ ИЭК.xlsx", workbook([
            ["№", "Код 1с", "Артикул поставщика", "Наименование", "Мин. разр. к отгр."],
            [1, "001_", "ARTICLE-01", "Кабель тестовый", 12],
        ])),
        ("Путь ИЭК 22.09.2026.xlsx", workbook([
            ["Код 1с", "Артикул ИЭК", "Наименование", "УТ-1 от 01.09.2026 (поступление до 01.10.2026)", "УТ-2 (поступление до 01.12.2026)"],
            ["001_", "ARTICLE-01", "Кабель тестовый", 3, 999],
        ])),
    )


def test_iek_import_to_real_engine_approval_and_export(iek_files, tmp_path):
    reports = inspect_files(iek_files)
    assert {r.kind for r in reports} == {"sales", "stock", "transit", "moq"}
    assert not missing_reports(reports)
    data, warnings = backend.load_partner_files(iek_files, settings(), "2026-09-23")
    assert len(data["sales"]) == 8 and data["sales"].qty.sum() == 800
    assert set(data["sales"].warehouse) == {WAREHOUSE}
    assert list(data["products"].sku) == ["001_"]
    assert data["products"].iloc[0].moq == 12
    assert data["products"].iloc[0].unit == "м"
    assert data["stock"].iloc[0].on_hand == 2
    assert data["in_transit"].eta.min() == pd.Timestamp("2026-10-01")
    assert data["stockouts"].empty and data["bom"].empty
    assert any("01.09.2026" in note for note in warnings)
    result = backend.calculate(data, {"today": "2026-09-23"})
    assert result.iloc[0].recommended_qty > 0
    assert result.iloc[0].in_transit == 3  # The December delivery is outside the horizon.
    assert result.iloc[0].oneoff_excluded == 0
    assert result.iloc[0].confidence == "low"
    assert "шт" not in result.iloc[0].explanation
    draft = orders.edited_order(result, result.assign(recommended_qty=24))
    approved = orders.approve(draft, "Тест", {}, tmp_path)
    exported = load_workbook(io.BytesIO(orders.xlsx_bytes(pd.DataFrame(approved["items"]))))
    assert exported.active["A2"].value == "ARTICLE-01"
    assert exported.active["D2"].value == 24
    assert exported.active["E2"].value == "м"
    csv = pd.read_csv(io.BytesIO(orders.csv_bytes(draft)), sep=";")
    assert csv.iloc[0]["Ед."] == "м"


def test_partial_transit_explains_missing_inputs(iek_files):
    files = iek_files[-1:]
    reports = inspect_files(files)
    assert reports[0].kind == "transit"
    assert len(missing_reports(reports)) == 2
    with pytest.raises(ValueError, match="добавьте.*продажи.*остатки"):
        convert_files(files, reports, settings(), "2026-09-23")


def test_movements_sign_aggregation_and_no_double_count(iek_files):
    movement = ("Динамика продаж_2025-2026.xlsx", workbook([
        ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"],
        ["02.08.2026 14:00:00", "1", "Расходная накладная", "001_", "Кабель", "м", "Алматы", -20],
        ["03.08.2026 14:00:00", "2", "Возврат от клиента", "001_", "Кабель", "м", "Астана", 3],
        ["03.09.2026 14:00:00", "3", "Расходная накладная", "001_", "Кабель", "м", "Алматы", -1000],
    ]))
    data, _ = backend.load_partner_files((*iek_files, movement), settings(), "2026-09-23")
    assert data["sales"].qty.sum() == 800
    data, _ = backend.load_partner_files((*iek_files[1:], movement), settings(), "2026-09-23")
    assert data["sales"].qty.sum() == 17 and len(data["sales"]) == 1
    assert data["sales"].date.iloc[0] == pd.Timestamp("2026-08-01")


def test_systeme_overview_and_moq():
    files = (
        ("Товар в пути SystemElectric 22.09.2026.xlsx", workbook([
            [None, "СКЛАДЫ"],
            ["Код 1с", "Артикул поставщика", "Наименование", "Категория 2026", "Январь 2026 г.", "Февраль 2026 г.", "Август 2026 г.", "Сентябрь 2026 г.", "Остаток", "Зарезервировано", "Свободный остаток", "СЭ в пути 24.09"],
            ["0002_", "SYS-2", "Розетка тест", 7, 10, 20, 30, 1000, 100, 60, 40, 5],
        ], "TDSheet")),
        ("MOQ SystemElectric.xlsx", workbook([
            ["№", "Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"],
            [], [1, "Розетка тест", "0002_", "SYS-2", 20],
        ])),
    )
    data, _ = backend.load_partner_files(files, settings(), "2026-09-23")
    assert data["sales"].qty.sum() == 60
    assert data["stock"].on_hand.iloc[0] == 40
    assert data["products"].pack_size.iloc[0] == 20
    assert data["products"].moq.iloc[0] == 0
    assert data["in_transit"].eta.iloc[0] == pd.Timestamp("2026-09-24")
    with pytest.raises(ValueError, match="позже"):
        backend.load_partner_files(files, settings(), "2026-09-21")


def test_numeric_codes_and_zero_history_months_survive_loader():
    files = (
        ("Продажи ИЭК.xlsx", workbook([
            ["Номенклатура", "Номенклатура.Код", "янв. 2026", "июль 2026", "авг. 2026"],
            ["Тест", "000123", 100, 0, 0],
        ])),
        ("Остатки ИЭК.xlsx", workbook([
            ["Номенклатура", "Номенклатура.Код", "Ед.", "сент. 2026"],
            ["Тест", "000123", "шт", 2],
        ])),
    )
    data, _ = backend.load_partner_files(files, settings(), "2026-09-23")
    assert data["products"].sku.iloc[0] == "000123"
    assert data["sales"].date.max() == pd.Timestamp("2026-08-01")
    assert list(data["sales"].qty) == [100, 0, 0]
    graph = backend.series(data, "000123", WAREHOUSE, "2026-09-23")
    assert graph.loc[graph.month.eq("2026-08"), "raw"].iloc[0] == 0


def test_import_rejects_unconfirmed_duplicate_and_bad_numbers(iek_files):
    with pytest.raises(ValueError, match="Подтвердите"):
        backend.load_partner_files(iek_files, settings(confirmed=False), "2026-09-23")
    with pytest.raises(ValueError, match="несколько отчётов"):
        inspect_files((*iek_files, iek_files[0]))
    broken = ("Остатки ИЭК.xlsx", workbook([
        ["Номенклатура", "Номенклатура.Код", "Ед.", "сент. 2026"],
        ["Тест", "001_", "м", "не число"],
    ]))
    files = (iek_files[0], broken)
    with pytest.raises(ValueError, match="нечисловое"):
        backend.load_partner_files(files, settings(), "2026-09-23")
    with pytest.raises(ValueError, match="не удалось открыть"):
        inspect_files((("Путь ИЭК.xlsx", b"not a zip"),))


def test_upload_page_accepts_partner_names_and_invalidates_settings(iek_files):
    app = AppTest.from_file(str(APP / "streamlit_app.py"), default_timeout=20)
    app.session_state.saved_source = "Загрузить файлы"
    app.session_state.saved_uploads = iek_files
    app.run()
    assert not app.exception
    assert any("Распознано" in el.value for el in app.success)
    next(box for box in app.checkbox if box.key == "partner_confirmed").check().run()
    next(btn for btn in app.button if btn.label == "Подготовить данные").click().run()
    next(btn for btn in app.button if btn.label == "Рассчитать").click().run()
    assert not app.exception
    assert not app.session_state.result.empty
    app.number_input(key="lead_IEK").set_value(20).run()
    assert next(btn for btn in app.button if btn.label == "Рассчитать").disabled
    assert all(btn.disabled for btn in app.button if btn.label.startswith("Утвердить заказ"))


def test_checks_cache_timestamp_and_strict_stockout(monkeypatch):
    from app import checks
    before = pd.DataFrame([dict(sku="DEMO-STOCKOUT", warehouse="WH", forecast_horizon=100, need_raw=200)])
    after = before.assign(forecast_horizon=150, need_raw=190)
    data = {"stockouts": pd.DataFrame({"sku": ["DEMO-STOCKOUT"]})}
    rows = checks.stockouts(data, {}, lambda d, p: after if len(d["stockouts"]) else before)
    assert rows[0]["Результат"] == "✅" and rows[1]["Результат"] == "❌"
    calls = []
    def run(*args):
        calls.append(1)
        return [checks.CheckResult("Тест", "passed", "ОК", elapsed_seconds=1.25)]
    monkeypatch.setattr(checks, "run_checks", run)
    app = AppTest.from_file(str(APP / "pages/2_Проверки.py"))
    app.session_state.calculation_data = {}
    app.session_state.calculation_params = {}
    app.session_state.calculation_signature = "first"
    app.run()
    next(b for b in app.button if b.label == "Прогнать все").click().run()
    next(b for b in app.button if b.label == "Прогнать все").click().run()
    assert len(calls) == 1 and not app.exception
    assert any("UTC" in c.value for c in app.caption)
    app.session_state.calculation_signature = "changed"
    app.run()
    assert not app.metric
    next(b for b in app.button if b.label == "Прогнать все").click().run()
    assert len(calls) == 2
